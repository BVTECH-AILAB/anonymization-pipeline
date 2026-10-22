"""Fail-closed verification of a redacted PDF (pypdfium2 and pypdf).

1. under every box no original text survives;
2. no text is left outside the visible page area;
3. no redacted value appears anywhere in the output text, decoded file
   objects or metadata, and no annotation, bookmark or document entry that can
   carry text is left;
4. preservation: rendered pages are identical outside the boxes.
"""

import re
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pypdfium2 as pdfium
import pypdfium2.raw as pdfium_c
from docx import Document
from pypdf import PdfReader
from pypdf.generic import IndirectObject, StreamObject

from anonymization_core.pdf_pages import (
    FLATTEN_DPI,
    page_box,
    pdfium_locked,
    render_page,
    to_pdf,
)
from anonymization_core.pipeline import Analysis, to_display
from anonymization_core.taxonomy import STRONG_LABELS


@dataclass
class Report:
    ok: bool = True
    issues: list[str] = field(default_factory=list)
    pages: set[int] = field(default_factory=set)  # 0-based pages with a failed check
    whole_file: bool = False  # a check failed on the file, not on one page
    pixel_changed_outside: list[float] = field(
        default_factory=list
    )  # per page, fraction of pixels

    def fail(self, msg: str, page: int | None = None) -> None:
        """Record a failed check: the document is not released.

        Args:
            msg: What failed, shown to the uploader.
            page: The 0-based page it failed on, or None for the whole file.
        """
        self.ok = False
        self.issues.append(msg if page is None else f"p{page + 1}: {msg}")
        if page is None:
            self.whole_file = True
        else:
            self.pages.add(page)


@dataclass
class Release:
    """Outcome of ``redact_and_verify``."""

    report: Report
    flattened: list[int] = field(
        default_factory=list
    )  # 0-based pages released as images
    seconds: dict[str, float] = field(default_factory=dict)  # redaction, verification


def _norm(s: str) -> str:
    return re.sub(r"\s+", "", s).lower()


def strong_values(analysis: Analysis) -> set[str]:
    """Redacted values that must not appear anywhere in an output: identifiers and full names."""
    boxes = [x for x in analysis.boxes if x.redact]
    values = {
        x.text for x in boxes if x.label in STRONG_LABELS and len(_norm(x.text)) >= 6
    }
    values |= {
        x.text for x in boxes if x.label == "PERSON" and len(x.text.split()) >= 2
    }
    # an occurrence the user chose to keep visible is not a leak
    return values - {
        x.text
        for x in analysis.boxes
        if not x.redact and "removed-by-user" in x.sources
    }


def verify_docx(path: Path, analysis: Analysis) -> list[str]:
    """Issues of the Word export: a redacted value left in its text.

    Returns:
        One message per value found; empty when the export is clean.
    """
    text = _norm(" ".join(p.text for p in Document(str(path)).paragraphs))
    return [
        f"value still in the Word export: '{v}'"
        for v in strong_values(analysis)
        if _norm(v) in text
    ]


def verify_markdown(path: Path, analysis: Analysis) -> list[str]:
    """Issues of the Markdown export: a redacted value left in its text.

    Returns:
        One message per value found; empty when the export is clean.
    """
    text = _norm(path.read_text(encoding="utf-8"))
    return [
        f"value still in the Markdown export: '{v}'"
        for v in strong_values(analysis)
        if _norm(v) in text
    ]


def _words(
    textpage: pdfium.PdfTextPage, box: tuple[float, float, float, float]
) -> list[str]:
    left, bottom, right, top = box
    return textpage.get_text_bounded(left, bottom, right, top).split()


def _file_text(path: Path) -> tuple[bytes, str, PdfReader]:
    """Every object of the file, streams decoded, and its metadata as text."""
    reader = PdfReader(path)
    chunks: list[bytes] = []
    for gen, ids in reader.xref.items():
        for idnum in ids:
            try:
                obj = reader.get_object(IndirectObject(idnum, gen, reader))
            except Exception:  # noqa: S112 - a broken object cannot hide a value from the others
                continue
            if isinstance(obj, StreamObject):
                try:
                    chunks.append(obj.get_data())
                except Exception:  # noqa: S112
                    chunks.append(obj._data)
            chunks.append(str(obj).encode("latin-1", "ignore"))
    for idnum in reader.xref_objStm:
        try:
            obj = reader.get_object(IndirectObject(idnum, 0, reader))
        except Exception:  # noqa: S112
            continue
        chunks.append(str(obj).encode("latin-1", "ignore"))
    meta = " ".join(str(v) for v in (reader.metadata or {}).values() if v)
    xmp = reader.xmp_metadata
    if xmp is not None:
        meta += " " + str(xmp.stream.get_data(), "utf-8", "ignore")
    return b"".join(chunks), meta, reader


@pdfium_locked
def verify(
    src: Path,
    out: Path,
    analysis: Analysis,
    flat: set[int],
    dpi: int = 72,
    tol: float = 2e-4,
) -> Report:
    """Fail-closed checks of ``out`` against ``src``: nothing left under the boxes, no
    value or text outside the page left in the file, pixels unchanged outside the boxes.

    Args:
        src: Original PDF.
        out: Redacted PDF.
        analysis: Boxes; ``applied`` holds the painted rectangles.
        flat: Pages replaced by a rendering (compared at ``FLATTEN_DPI``).
        dpi: Resolution of the pixel comparison of the other pages.
        tol: Share of pixels allowed to change outside the boxes.
    """
    rep = Report()
    boxes = [x for x in analysis.boxes if x.redact]
    a, b = pdfium.PdfDocument(src), pdfium.PdfDocument(out)
    try:
        if len(a) != len(b):
            rep.fail(f"page count changed: {len(a)} -> {len(b)}")
            return rep
        texts: list[str] = []
        for pno in range(len(b)):
            page = b[pno]
            textpage = page.get_textpage()
            texts.append(textpage.get_text_bounded())
            # 1. nothing left under the boxes
            for x in boxes:
                if x.page != pno or x.layer == "visual":
                    continue
                # a flattened page is stored upright, as displayed
                r = to_display(x.rect, analysis.pages[pno]) if pno in flat else x.rect
                inner = (
                    r[0] + 1,
                    r[1] + (r[3] - r[1]) * 0.3,
                    r[2] - 1,
                    r[3] - (r[3] - r[1]) * 0.3,
                )
                if inner[2] <= inner[0] or inner[3] <= inner[1]:
                    continue
                orig = _norm(x.text)
                for w in _words(textpage, to_pdf(page, inner)):
                    if len(w) >= 2 and _norm(w) in orig:
                        rep.fail(f"'{w}' still under {x.label} box", pno)
            # 1b. text outside the visible page area
            cx0, cy0, cx1, cy1 = page_box(page)
            outside = 0
            for i in range(textpage.count_chars()):
                left, bottom, right, top = textpage.get_charbox(i)
                if right < cx0 or left > cx1 or top < cy0 or bottom > cy1:
                    outside += 1
            if outside:
                rep.fail(f"{outside} characters outside the visible page area", pno)
            # annotations
            if pdfium_c.FPDFPage_GetAnnotCount(page.raw) > 0:
                rep.fail("annotations left", pno)

        # 2. global search for strong values
        full = _norm(" ".join(texts))
        raw, meta, reader = _file_text(out)
        raw_low = raw.lower()
        for v in strong_values(analysis):
            nv = _norm(v)
            if nv and nv in full:
                rep.fail(f"value still in text: '{v}'")
            if len(v) >= 6 and v.lower().encode("latin-1", "ignore") in raw_low:
                rep.fail(f"value still in file bytes: '{v}'")
            if nv and nv in _norm(meta):
                rep.fail(f"value still in metadata: '{v}'")
        if reader.outline:
            rep.fail("bookmarks left")
        root = reader.root_object
        for key in ("/StructTreeRoot", "/Names", "/AcroForm", "/OpenAction"):
            if key in root:
                rep.fail(f"document entry left: {key[1:]}")

        # 3. preservation outside the boxes
        for pno in range(len(a)):
            scale = (FLATTEN_DPI if pno in flat else dpi) / 72
            A = render_page(a[pno], scale).astype(np.int16)
            B = render_page(b[pno], scale).astype(np.int16)
            if A.shape != B.shape:
                rep.fail(f"page size changed: {A.shape} -> {B.shape}", pno)
                continue
            diff = np.abs(A - B).max(axis=2) > 40
            info = analysis.pages[pno]
            rects = [r for n, r, _ in analysis.applied if n == pno] or [
                x.rect for x in boxes if x.page == pno
            ]
            for rect in rects:
                x0, y0, x1, y1 = to_display(rect, info)
                diff[
                    max(0, int((y0 - 3) * scale)) : int((y1 + 3) * scale) + 1,
                    max(0, int((x0 - 3) * scale)) : int((x1 + 3) * scale) + 1,
                ] = False
            frac = float(diff.mean())
            rep.pixel_changed_outside.append(frac)
            if frac > tol:
                rep.fail(f"{frac:.2%} of pixels changed outside the boxes", pno)
        return rep
    finally:
        a.close()
        b.close()
