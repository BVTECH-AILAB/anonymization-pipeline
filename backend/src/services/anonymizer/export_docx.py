"""Word export: the anonymized text, line by line, with redacted parts as black blocks.
Layout is preserved only in the PDF; this is for editing the text."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor

from src.services.anonymizer.geometry import (
    Rect,
    contains_point,
    grow,
)
from src.services.anonymizer.layers import Token
from src.services.anonymizer.pipeline import Analysis

# (area, label) of a redaction box, grown by half a point
Redaction = tuple[Rect, str]


def redactions(analysis: Analysis) -> dict[int, list[Redaction]]:
    """Redacted areas of every page, with their label."""
    red: dict[int, list[Redaction]] = {}
    for b in analysis.boxes:
        if b.redact:
            red.setdefault(b.page, []).append((grow(b.rect, 0.5), b.label))
    return red


def page_lines(analysis: Analysis, pno: int) -> list[list[Token]]:
    """Tokens of one page grouped by line, from the layer the text exports read.

    A scanned page is read from its OCR layer, any other page from the native
    one; tokens read with a confidence under 0.2 are left out.
    """
    layers = [L for L in analysis.layers if L.page == pno]
    layer = (
        next((L for L in layers if L.source == "ocr"), None)
        if analysis.pages[pno].get("scanned")
        else None
    )
    layer = layer or next((L for L in layers if L.source == "native"), None)
    if layer is None:
        return []
    lines: dict[int, list[Token]] = {}
    for t in layer.tokens:
        if t.conf < 0.2:
            continue
        lines.setdefault(t.line, []).append(t)
    return list(lines.values())


def char_label(red: list[Redaction], char: Rect) -> str | None:
    """Label of the redaction covering the centre of ``char``, None when it stays visible."""
    cx, cy = (char[0] + char[2]) / 2, (char[1] + char[3]) / 2
    return next((label for r, label in red if contains_point(r, cx, cy)), None)


def export_docx(analysis: Analysis, out: Path) -> None:
    """Write the anonymized text of ``analysis`` to ``out``; redacted parts become black blocks."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(10)
    doc.core_properties.author = ""
    red = redactions(analysis)
    for pno in range(len(analysis.pages)):
        if pno:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        rects = red.get(pno, [])
        for toks in page_lines(analysis, pno):
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(2)
            for i, t in enumerate(toks):
                if i:
                    para.add_run(" ")
                buf, hidden = "", None
                for ch, cr in zip(t.text, t.chars, strict=False):
                    h = char_label(rects, cr) is not None
                    if hidden is not None and h != hidden:
                        _run(para, buf, hidden)
                        buf = ""
                    buf += "█" if h else ch
                    hidden = h
                _run(para, buf, hidden)
    doc.save(str(out))


def _run(para, text: str, hidden: bool | None) -> None:
    if not text:
        return
    r = para.add_run(text)
    if hidden:
        r.font.color.rgb = RGBColor(0, 0, 0)
