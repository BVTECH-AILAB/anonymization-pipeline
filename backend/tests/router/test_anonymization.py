"""The anonymization API, shaped like the document agent of ai-platform-v2.

A session is created by the upload; its work runs inside an SSE stream whose
frames are ``status`` (progress), ``done`` (the result) or ``error``.
"""

import io
import json
import shutil
import time
from pathlib import Path

import pypdfium2 as pdfium
import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from src.errors import (
    FileTooLargeError,
    ServiceUnavailableError,
    TooManyJobsError,
    TooManyPagesError,
)
from src.router.anonymization import open_file_response
from src.services.anonymization import JobManager
from src.services.anonymizer.errors import NerUnavailableError
from src.services.anonymizer.pipeline import Pipeline
from src.services.anonymizer.reader import LiteparseReader
from src.services.anonymizer.verify import Report, verify
from tests.conftest import FAKE_CF
from tests.documents import (
    PATIENT,
    PHONE,
    FakeReader,
    draw_pdf,
    edit_pdf,
    letter_pdf,
    native_words,
    scanned_pdf,
)
from tests.sessions import BASE, anonymize, create, edit, frames, run


def pdf_text(data: bytes) -> str:
    with pdfium.PdfDocument(data) as doc:
        return " ".join(
            doc[i].get_textpage().get_text_bounded() for i in range(len(doc))
        )


# --- sessions and the anonymization stream ----------------------------------


def test_labels(client: TestClient) -> None:
    labels = {lab["code"]: lab for lab in client.get("/anonymizations/labels").json()}

    assert labels["FISCAL_CODE"] == {
        "code": "FISCAL_CODE",
        "title": "Codice fiscale",
        "kind": "text",
        "redact": True,
    }
    assert labels["SIGNATURE"]["kind"] == "visual"


def test_a_new_session_waits_for_its_stream(client: TestClient, letter: Path) -> None:
    session = create(client, letter)

    assert session["status"] == "created"
    assert session["mode"] == "accurate"
    assert session["pages"] == 1
    assert session["result"] is None
    assert client.get(f"{BASE}/{session['session_id']}").json()["status"] == "created"


def test_the_stream_reports_progress_then_the_result(
    client: TestClient, letter: Path
) -> None:
    session_id = create(client, letter)["session_id"]

    stream = anonymize(client, session_id)

    events = [event for event, _ in stream]
    assert events[-1] == "done"
    assert set(events[:-1]) == {"status"}
    percents = [data["percent"] for event, data in stream if event == "status"]
    assert percents == sorted(percents)
    assert all(data["text"] for event, data in stream if event == "status")
    done = stream[-1][1]
    assert done["elapsed_ms"] >= 0
    session = done["session"]
    assert session["session_id"] == session_id
    assert session["status"] == "done"
    result = session["result"]
    assert result["verified"] is True
    assert result["counts"]["FISCAL_CODE"] == 1
    assert result["counts"]["PERSON"] >= 2
    assert result["counts"]["QRCODE"] == 1
    assert result["counts"]["SIGNATURE"] == 1
    assert set(result["timings"]) == {
        "text_layers",
        "detection",
        "visual",
        "redaction",
        "verification",
    }


def test_the_stream_of_a_finished_session_ends_at_once(
    client: TestClient, letter: Path
) -> None:
    session_id = create(client, letter)["session_id"]
    anonymize(client, session_id)

    again = anonymize(client, session_id)

    assert [event for event, _ in again] == ["done"]


def test_a_second_stream_while_running_is_busy(
    client: TestClient, jobs: JobManager, letter: Path
) -> None:
    session_id = create(client, letter)["session_id"]
    jobs.get(session_id).status = "running"

    stream = anonymize(client, session_id)

    assert stream == [
        (
            "error",
            {
                "error_title": "Elaborazione in corso",
                "error_message": "Il documento è già in elaborazione: attendi che finisca.",
                "error_code": "SESSION_BUSY",
            },
        )
    ]


def test_an_unknown_session_is_not_found(client: TestClient) -> None:
    response = client.post(f"{BASE}/{'0' * 32}/anonymize/stream")

    assert response.status_code == 404
    assert response.json()["code"] == "JOB_NOT_FOUND"


def test_processing_error_is_an_error_frame(
    client: TestClient, jobs: JobManager, letter: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def broken(path: Path, progress, visual: bool = True) -> object:
        raise RuntimeError(f"cannot read {FAKE_CF}")

    monkeypatch.setattr(jobs.pipeline, "analyze", broken)
    session_id = create(client, letter)["session_id"]

    stream = anonymize(client, session_id)

    event, data = stream[-1]
    assert event == "error"
    assert data["error_code"] == "PROCESSING_FAILED"
    assert data["error_title"] and data["error_message"]
    assert FAKE_CF not in json.dumps(stream)
    state = client.get(f"{BASE}/{session_id}").json()
    assert (state["status"], state["error_code"]) == ("error", "PROCESSING_FAILED")


def test_a_dependency_down_is_reported_with_its_code(
    client: TestClient, jobs: JobManager, letter: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def ner_down(path: Path, progress, visual: bool = True) -> object:
        raise NerUnavailableError()

    monkeypatch.setattr(jobs.pipeline, "analyze", ner_down)
    session_id = create(client, letter, mode="fast")["session_id"]

    event, data = anonymize(client, session_id)[-1]

    assert (event, data["error_code"]) == ("error", "NER_UNAVAILABLE")
    assert data["error_title"] == "Servizio non disponibile"


def test_cancel(
    client: TestClient, jobs: JobManager, letter: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    original = jobs.pipeline.analyze

    def slow_analyze(path: Path, progress, visual: bool = True) -> object:
        progress("read", 0)
        time.sleep(0.3)
        progress("read", 0.5)
        return original(path, progress, visual)

    monkeypatch.setattr(jobs.pipeline, "analyze", slow_analyze)
    session_id = create(client, letter)["session_id"]
    jobs.start(session_id)

    assert client.post(f"{BASE}/{session_id}/cancel").status_code == 202
    deadline = time.monotonic() + 10
    while session_id in jobs.jobs and time.monotonic() < deadline:
        time.sleep(0.05)

    assert session_id not in jobs.jobs
    assert client.get(f"{BASE}/{session_id}").status_code == 404


def test_cancel_before_the_stream_deletes_the_session(
    client: TestClient, jobs: JobManager, letter: Path
) -> None:
    session_id = create(client, letter)["session_id"]
    job_dir = jobs.get(session_id).dir

    assert client.post(f"{BASE}/{session_id}/cancel").status_code == 202

    assert not job_dir.exists()
    assert client.get(f"{BASE}/{session_id}").status_code == 404


def test_delete_removes_files(
    client: TestClient, jobs: JobManager, letter: Path
) -> None:
    session_id = run(client, letter)["session_id"]
    job_dir = jobs.get(session_id).dir

    assert client.delete(f"{BASE}/{session_id}").status_code == 204
    assert not job_dir.exists()
    assert client.get(f"{BASE}/{session_id}").status_code == 404


def test_without_a_job_manager_the_api_is_unavailable(client: TestClient) -> None:
    client.app.state.jobs = None

    response = client.get(f"{BASE}/{'0' * 32}")

    assert response.status_code == 503


# --- uploads ----------------------------------------------------------------


@pytest.mark.parametrize(
    ("name", "data", "status", "code"),
    [
        ("a.txt", b"hello", 415, "UNSUPPORTED_FORMAT"),
        ("a.pdf", b"not a pdf", 422, "UNREADABLE_DOCUMENT"),
        ("a.png", b"not an image", 422, "UNREADABLE_DOCUMENT"),
    ],
)
def test_upload_errors(
    client: TestClient, name: str, data: bytes, status: int, code: str
) -> None:
    response = client.post(
        BASE, files={"file": (name, data, "application/octet-stream")}
    )

    assert response.status_code == status
    assert response.json()["code"] == code


def test_password_protected_pdf(client: TestClient, tmp_path: Path) -> None:
    path = draw_pdf(tmp_path / "locked.pdf", lambda p: p.text(72, 72, "segreto"))
    edit_pdf(path, password="pw")

    response = client.post(
        BASE, files={"file": ("locked.pdf", path.read_bytes(), "application/pdf")}
    )

    assert response.status_code == 422
    assert response.json()["code"] == "PASSWORD_PROTECTED"


def test_limits(jobs: JobManager, letter: Path, tmp_path: Path) -> None:
    jobs.settings = jobs.settings.model_copy(update={"max_upload_mb": 1})
    with pytest.raises(FileTooLargeError):
        jobs.create(b"x" * (1024 * 1024 + 1), "big.pdf")

    jobs.settings = jobs.settings.model_copy(
        update={"max_upload_mb": 30, "max_pages": 1}
    )
    many = draw_pdf(tmp_path / "many.pdf", lambda p: None, lambda p: None)
    with pytest.raises(TooManyPagesError):
        jobs.create(many.read_bytes(), "many.pdf")

    jobs.settings = jobs.settings.model_copy(update={"max_pages": 50, "max_jobs": 1})
    jobs.create(letter.read_bytes(), "a.pdf")
    with pytest.raises(TooManyJobsError):
        jobs.create(letter.read_bytes(), "b.pdf")


def test_expired_jobs_are_purged(jobs: JobManager, letter: Path) -> None:
    job = jobs.create(letter.read_bytes(), "a.pdf")
    job.created -= jobs.settings.job_ttl_minutes * 60 + 1

    assert jobs.purge_expired() == 1
    assert job.id not in jobs.jobs
    assert not job.dir.exists()


def test_image_upload_becomes_a_pdf(
    client: TestClient, letter: Path, tmp_path: Path
) -> None:
    with pdfium.PdfDocument(letter) as doc:
        img = doc[0].render(scale=80 / 72).to_pil().convert("RGB")
    photo = tmp_path / "foto.png"
    img.save(photo, "PNG")

    state = run(client, photo, mime="image/png")

    assert state["status"] == "done"
    assert state["result"]["pages"][0]["scanned"] is True


def test_image_with_too_many_frames(jobs: JobManager) -> None:
    jobs.settings = jobs.settings.model_copy(update={"max_pages": 1})
    frames_ = [Image.new("RGB", (40, 40), c) for c in ("white", "black")]
    buf = io.BytesIO()
    frames_[0].save(buf, "TIFF", save_all=True, append_images=frames_[1:])

    with pytest.raises(TooManyPagesError):
        jobs.create(buf.getvalue(), "scan.tiff")


def test_letter_fixture_contains_the_pii(letter: Path) -> None:
    text = pdf_text(letter.read_bytes())

    assert FAKE_CF in text
    assert PATIENT in text


def test_scanned_upload_uses_ocr(
    client: TestClient, fake_reader: FakeReader, tmp_path: Path
) -> None:
    native = letter_pdf(tmp_path / "native.pdf")
    scan = scanned_pdf(native, tmp_path / "scan.pdf")
    fake_reader.words = native_words(native)

    result = run(client, scan)["result"]

    assert result["verified"] is True
    assert result["pages"][0]["scanned"] is True
    assert result["counts"]["FISCAL_CODE"] == 1
    assert all(b["layer"] in ("ocr", "visual") for b in result["boxes"])


def test_upload_refused_while_models_load(
    settings, tmp_path: Path, letter: Path
) -> None:
    loading = JobManager(
        Pipeline(LiteparseReader(""), None),  # no OCR server: never ready
        settings.model_copy(update={"preload_models": True}),
        tmp_path / "w",
    )
    try:
        with pytest.raises(ServiceUnavailableError):
            loading.create(letter.read_bytes(), "a.pdf")
    finally:
        loading.shutdown()


def test_mode_defaults_to_accurate(client: TestClient, letter: Path) -> None:
    with letter.open("rb") as f:
        created = client.post(
            BASE, files={"file": ("a.pdf", f, "application/pdf")}
        ).json()

    assert created["mode"] == "accurate"
    assert client.get(f"{BASE}/{created['session_id']}").json()["mode"] == "accurate"


def test_unknown_mode_is_refused(client: TestClient, letter: Path) -> None:
    with letter.open("rb") as f:
        response = client.post(
            BASE,
            files={"file": ("a.pdf", f, "application/pdf")},
            data={"mode": "turbo"},
        )

    assert response.status_code == 422


# --- downloads and preview --------------------------------------------------


def test_download_pdf_without_pii_and_with_generic_name(
    client: TestClient, letter: Path
) -> None:
    session_id = run(client, letter, name=f"referto_{PATIENT}_{FAKE_CF}.pdf")[
        "session_id"
    ]

    response = client.get(f"{BASE}/{session_id}/download")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    disposition = response.headers["content-disposition"]
    assert disposition == "attachment; filename*=UTF-8''documento_anonimizzato.pdf"
    assert response.headers["cache-control"] == "no-store"
    text = pdf_text(response.content)
    for value in (FAKE_CF, PATIENT, PHONE, "Rossi"):
        assert value not in text
    assert "morbo di Crohn" in text
    assert "Babinski" in text
    assert "LETTERA DI DIMISSIONE" in text


def test_download_docx(client: TestClient, letter: Path) -> None:
    session_id = run(client, letter)["session_id"]

    response = client.get(f"{BASE}/{session_id}/download?format=docx")

    assert response.status_code == 200
    text = " ".join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)
    assert FAKE_CF not in text
    assert "█" in text
    assert "Crohn" in text


def test_preview_is_the_released_pdf_inline(client: TestClient, letter: Path) -> None:
    session_id = run(client, letter)["session_id"]

    response = client.get(f"{BASE}/{session_id}/preview")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].startswith("inline;")
    assert PATIENT not in pdf_text(response.content)


def test_pages_with_boxes_are_released_as_images(
    client: TestClient, letter: Path
) -> None:
    state = run(client, letter)

    assert state["result"]["verified"] is True
    assert state["result"]["flattened_pages"] == [1]  # 1-based, as shown to the user
    pdf = client.get(f"{BASE}/{state['session_id']}/download").content
    assert PATIENT not in pdf_text(pdf)
    with pdfium.PdfDocument(pdf) as doc:  # rendered page: still a document, not blank
        assert doc[0].render(scale=1).to_numpy().mean() < 255


@pytest.mark.parametrize("path", ["download", "preview"])
def test_release_refused_when_a_check_fails(
    client: TestClient, letter: Path, monkeypatch: pytest.MonkeyPatch, path: str
) -> None:
    def failing(*args, **kwargs) -> Report:
        rep = Report()
        rep.fail("value still in text: 'x'")
        return rep

    monkeypatch.setattr("src.services.anonymizer.release.verify", failing)
    state = run(client, letter)

    assert state["result"]["verified"] is False
    assert state["result"]["issues"]
    response = client.get(f"{BASE}/{state['session_id']}/{path}")
    assert response.status_code == 409
    assert response.json()["code"] == "VERIFICATION_FAILED"


def test_verification_time_is_measured(
    client: TestClient, letter: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def slow_verify(*args, **kwargs):
        time.sleep(0.05)
        return verify(*args, **kwargs)

    monkeypatch.setattr("src.services.anonymizer.release.verify", slow_verify)
    timings = run(client, letter)["result"]["timings"]

    assert timings["verification"] >= 0.05


def test_release_before_completion_is_refused(client: TestClient, letter: Path) -> None:
    session_id = create(client, letter)["session_id"]

    response = client.get(f"{BASE}/{session_id}/download")

    assert response.status_code == 409
    assert response.json()["code"] == "JOB_NOT_READY"


def test_download_during_an_edit_is_not_ready(
    client: TestClient, jobs: JobManager, letter: Path
) -> None:
    session_id = run(client, letter)["session_id"]
    jobs.get(session_id).pending_edits = 1

    response = client.get(f"{BASE}/{session_id}/download")

    assert response.status_code == 409
    assert response.json()["code"] == "JOB_NOT_READY"


# --- fast mode ---------------------------------------------------------------


def test_fast_mode_releases_anonymized_markdown(
    client: TestClient, letter: Path
) -> None:
    state = run(client, letter, name=f"referto_{FAKE_CF}.pdf", mode="fast")

    assert state["mode"] == "fast"
    result = state["result"]
    assert result["verified"] is True
    assert result["counts"]["FISCAL_CODE"] == 1
    assert "boxes" not in result
    assert "QRCODE" not in result["counts"]  # no visual detection in fast mode
    response = client.get(f"{BASE}/{state['session_id']}/download")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/markdown")
    assert "documento_anonimizzato.md" in response.headers["content-disposition"]
    md = response.text
    assert md.startswith("<!-- pagina 1 -->")
    assert "[CODICE FISCALE]" in md
    assert "[IMMAGINE]" in md
    for value in (FAKE_CF, PATIENT, PHONE):
        assert value not in md
    assert "Crohn" in md
    preview = client.get(f"{BASE}/{state['session_id']}/preview")
    assert preview.headers["content-disposition"].startswith("inline;")
    assert preview.text == md


def test_fast_markdown_refused_when_a_value_survives(
    client: TestClient, letter: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        "src.services.anonymizer.document.verify_markdown",
        lambda *_: ["value still in the Markdown export: 'x'"],
    )
    state = run(client, letter, mode="fast")

    assert state["result"]["verified"] is False
    response = client.get(f"{BASE}/{state['session_id']}/download?format=md")
    assert response.status_code == 409
    assert response.json()["code"] == "VERIFICATION_FAILED"


@pytest.mark.parametrize(
    ("mode", "path"),
    [
        ("fast", "download?format=pdf"),
        ("fast", "download?format=docx"),
        ("fast", "review"),
        ("fast", "pages/anonymized/0"),
        ("accurate", "download?format=md"),
    ],
)
def test_actions_of_the_other_mode_are_refused(
    client: TestClient, letter: Path, mode: str, path: str
) -> None:
    session_id = run(client, letter, mode=mode)["session_id"]

    response = client.get(f"{BASE}/{session_id}/{path}")

    assert response.status_code == 409
    assert response.json()["code"] == "MODE_NOT_SUPPORTED"


def test_edits_of_a_fast_session_are_refused(client: TestClient, letter: Path) -> None:
    session_id = run(client, letter, mode="fast")["session_id"]

    response = client.post(f"{BASE}/{session_id}/edits/stream", json={})

    assert response.status_code == 409
    assert response.json()["code"] == "MODE_NOT_SUPPORTED"


def test_fast_session_keeps_the_original_page(client: TestClient, letter: Path) -> None:
    session_id = run(client, letter, mode="fast")["session_id"]

    response = client.get(f"{BASE}/{session_id}/pages/original/0?dpi=60")

    assert response.status_code == 200


# --- review and edits ----------------------------------------------------------


def test_page_images(client: TestClient, letter: Path) -> None:
    session_id = run(client, letter)["session_id"]

    for which in ("original", "anonymized"):
        response = client.get(f"{BASE}/{session_id}/pages/{which}/0?dpi=60")
        assert response.status_code == 200
        assert response.headers["content-type"] == "image/png"
    assert client.get(f"{BASE}/{session_id}/pages/original/3").status_code == 404


def test_review_manifest(client: TestClient, letter: Path) -> None:
    session_id = run(client, letter)["session_id"]
    url = f"{BASE}/{session_id}/review"

    review = client.get(url).json()

    assert review["session_id"] == session_id
    assert review["mode"] == "accurate"
    assert review["version"] == 0
    assert review["verified"] is True
    page = review["pages"][0]
    assert page["index"] == 0
    assert (page["width"], page["height"]) == (595, 842)
    assert page["original"] == f"{BASE}/{session_id}/pages/original/0"
    assert page["anonymized"] == f"{BASE}/{session_id}/pages/anonymized/0?v=0"
    assert client.get(page["original"]).status_code == 200
    assert client.get(page["anonymized"]).status_code == 200
    box = review["boxes"][0]
    assert set(box) == {"id", "page", "rect", "label", "layer", "redact"}
    assert review["labels"]["PERSON"] == "Nomi"
    assert FAKE_CF not in json.dumps(review)
    assert PATIENT not in json.dumps(review)

    edit(client, session_id, {"remove": [0]})
    edited = client.get(url).json()

    assert edited["version"] == 1
    assert edited["pages"][0]["anonymized"].endswith("?v=1")
    assert edited["boxes"][0]["redact"] is False


def test_review_before_completion_is_refused(client: TestClient, letter: Path) -> None:
    session_id = create(client, letter)["session_id"]

    response = client.get(f"{BASE}/{session_id}/review")

    assert response.status_code == 409
    assert response.json()["code"] == "JOB_NOT_READY"


def test_edits_add_remove_and_search(client: TestClient, letter: Path) -> None:
    state = run(client, letter)
    session_id = state["session_id"]
    person = next(
        b for b in state["result"]["boxes"] if b["label"] == "PERSON" and b["redact"]
    )

    stream = edit(
        client,
        session_id,
        {
            "add": [{"page": 0, "rect": [40, 300, 200, 320]}],
            "remove": [person["id"]],
            "search": "Medicina Interna",
        },
    )

    assert {event for event, _ in stream[:-1]} <= {"status"}
    event, done = stream[-1]
    assert event == "done"
    assert done["search_hits"] == 1
    assert done["edit_seconds"] >= 0
    result = done["session"]["result"]
    assert result["edits"] == 1
    boxes = result["boxes"]
    assert boxes[person["id"]]["redact"] is False
    assert "removed-by-user" in boxes[person["id"]]["sources"]
    assert any(b["sources"] == ["manual-area"] for b in boxes)
    text = pdf_text(client.get(f"{BASE}/{session_id}/download").content)
    assert "Medicina Interna" not in text

    restored = edit(client, session_id, {"restore": [person["id"]]})[-1][1]
    assert restored["session"]["result"]["boxes"][person["id"]]["redact"] is True


def test_edits_before_completion_are_refused(client: TestClient, letter: Path) -> None:
    session_id = create(client, letter)["session_id"]

    response = client.post(f"{BASE}/{session_id}/edits/stream", json={"search": "x"})

    assert response.status_code == 409
    assert response.json()["code"] == "JOB_NOT_READY"


def test_a_second_edit_while_one_runs_is_busy(
    client: TestClient, jobs: JobManager, letter: Path
) -> None:
    session_id = run(client, letter)["session_id"]
    jobs.get(session_id).pending_edits = 1

    stream = edit(client, session_id, {"remove": [0]})

    assert [(e, d["error_code"]) for e, d in stream] == [("error", "SESSION_BUSY")]


def test_a_failed_edit_leaves_the_session_blocked_and_clean(
    client: TestClient, jobs: JobManager, letter: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    session_id = run(client, letter)["session_id"]

    def broken(*args, **kwargs):
        raise RuntimeError(f"redaction failed on {FAKE_CF}")

    monkeypatch.setattr("src.services.anonymizer.document.redact_and_verify", broken)
    stream = edit(client, session_id, {"remove": [0]})

    event, data = stream[-1]
    assert (event, data["error_code"]) == ("error", "PROCESSING_FAILED")
    assert FAKE_CF not in json.dumps(stream)
    state = client.get(f"{BASE}/{session_id}").json()
    assert state["result"]["verified"] is False
    assert state["result"]["edits"] == 0
    assert client.get(f"{BASE}/{session_id}/download").status_code == 409
    assert not list(jobs.get(session_id).dir.glob("anonymized.next.*"))
    assert jobs.get(session_id).pending_edits == 0


# --- files that disappear --------------------------------------------------------


@pytest.mark.parametrize(
    ("mode", "path"),
    [
        ("accurate", "download"),
        ("accurate", "download?format=docx"),
        ("accurate", "preview"),
        ("fast", "download"),
        ("accurate", "pages/original/0"),
        ("accurate", "pages/anonymized/0"),
    ],
)
def test_files_deleted_under_a_request_are_not_found(
    client: TestClient, jobs: JobManager, letter: Path, mode: str, path: str
) -> None:
    session_id = run(client, letter, mode=mode)["session_id"]
    shutil.rmtree(jobs.get(session_id).dir)  # expired between lookup and read

    response = client.get(f"{BASE}/{session_id}/{path}")

    assert response.status_code == 404
    assert response.json()["code"] == "JOB_NOT_FOUND"


async def test_a_file_deleted_while_it_is_sent_is_sent_whole(tmp_path: Path) -> None:
    path = tmp_path / "anonymized.pdf"
    data = bytes(range(256)) * 1000  # several chunks
    path.write_bytes(data)

    response = open_file_response(path, "application/pdf", "documento.pdf")
    path.unlink()  # the job expires after the response was built
    chunks = [chunk async for chunk in response.body_iterator]
    body = b"".join(c for c in chunks if isinstance(c, bytes))

    assert all(isinstance(c, bytes) for c in chunks)
    assert body == data
    assert response.headers["content-length"] == str(len(data))
    assert (
        response.headers["content-disposition"]
        == "attachment; filename*=UTF-8''documento.pdf"
    )


def test_frames_parser_skips_keep_alive() -> None:
    body = ': keep-alive\n\nevent: status\ndata: {"percent": 1}\n\n'

    assert frames(body) == [("status", {"percent": 1})]
