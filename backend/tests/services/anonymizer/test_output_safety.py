"""Output safety: rotated pages, text outside the page, the Word export."""

from pathlib import Path

import numpy as np
import pypdfium2 as pdfium
from docx import Document
from pypdf import PdfReader

from src.services.anonymizer.layers import Word, char_boxes
from src.services.anonymizer.pipeline import (
    Analysis,
    Box,
    Pipeline,
    from_display,
    to_display,
)
from src.services.anonymizer.release import redact_and_verify
from src.services.anonymizer.verify import verify_docx
from tests.documents import (
    FAKE_CF,
    FakeNER,
    FakeReader,
    draw_pdf,
    edit_pdf,
    find_text,
    scanned_pdf,
)


def _rotated_scan(tmp_path: Path) -> tuple[Path, tuple[float, float, float, float]]:
    """A scanned page stored with /Rotate 90 and the page-space box of the CF on it."""
    native = draw_pdf(
        tmp_path / "native.pdf", lambda p: p.text(100, 700, f"C.F. {FAKE_CF}", size=14)
    )
    cf = find_text(native, FAKE_CF)
    scan = scanned_pdf(native, tmp_path / "scan.pdf", dpi=100)
    return edit_pdf(scan, tmp_path / "rotated_scan.pdf", rotate={0: 90}), cf


def _dark_box(img: np.ndarray, scale: float) -> tuple[float, float, float, float]:
    ys, xs = np.nonzero(img.mean(axis=2) < 60)
    return (
        xs.min() / scale,
        ys.min() / scale,
        (xs.max() + 1) / scale,
        (ys.max() + 1) / scale,
    )


def test_ocr_boxes_of_a_rotated_scan_land_on_the_text(tmp_path: Path) -> None:
    path, cf = _rotated_scan(tmp_path)
    info = {
        "width": 842,
        "height": 595,
        "rotation": 90,
        "page_width": 595,
        "page_height": 842,
    }
    shown = to_display(cf, info)  # where the OCR sees the code on the rendering
    word = Word(FAKE_CF, shown, 0.99, "ocr", 0)
    analysis = Pipeline(FakeReader({0: [word]}), FakeNER(())).analyze(path)

    boxes = [b for b in analysis.boxes if b.label == "FISCAL_CODE"]
    assert boxes
    x0, y0, x1, y1 = boxes[0].rect
    assert x0 <= cf[0] + 1 and y0 <= cf[1] + 1 and x1 >= cf[2] - 1 and y1 >= cf[3] - 1

    out = tmp_path / "out.pdf"
    assert redact_and_verify(path, analysis, out).report.ok
    with pdfium.PdfDocument(out) as doc:
        img = doc[0].render(scale=1, rev_byteorder=True).to_numpy()[:, :, :3]
    sx0, sy0, sx1, sy1 = (int(v) for v in shown)
    assert img[sy0 + 2 : sy1 - 2, sx0 + 2 : sx1 - 2].mean() < 60  # under a black box


def test_character_boxes_follow_the_reading_direction() -> None:
    down = char_boxes((0, 0, 10, 30), 3, "ttb")
    up = char_boxes((0, 0, 10, 30), 3, "btt")

    assert down[0][1] == 0 and down[-1][3] == 30
    assert up[0][3] == 30 and up[-1][1] == 0


def test_text_outside_the_visible_page_is_not_released(tmp_path: Path) -> None:
    def draw(p) -> None:
        p.text(100, 100, "Referto", size=12)
        p.text(100, 900, f"C.F. {FAKE_CF}", size=12)  # below the page

    src = draw_pdf(tmp_path / "offpage.pdf", draw)
    analysis = Pipeline(FakeReader(), FakeNER(())).analyze(src)
    out = tmp_path / "out.pdf"

    release = redact_and_verify(src, analysis, out)

    assert release.report.ok, release.report.issues
    assert FAKE_CF.encode() not in out.read_bytes()
    with pdfium.PdfDocument(out) as doc:
        assert FAKE_CF not in doc[0].get_textpage().get_text_bounded()


def test_a_value_left_in_the_word_export_is_reported(tmp_path: Path) -> None:
    analysis = Analysis(tmp_path / "x.pdf", [{"width": 595, "height": 842}])
    analysis.boxes.append(
        Box(0, (10, 10, 90, 20), "PERSON", "Mario Rossi", 0.9, ["ner"], "native", True)
    )
    leaky, clean = tmp_path / "leaky.docx", tmp_path / "clean.docx"
    for path, text in ((leaky, "Paziente Mario Rossi"), (clean, "Paziente ███")):
        d = Document()
        d.add_paragraph(text)
        d.save(str(path))

    assert verify_docx(leaky, analysis) == [
        "value still in the Word export: 'Mario Rossi'"
    ]
    assert verify_docx(clean, analysis) == []


def test_display_space_matches_the_rendering_for_every_rotation(tmp_path: Path) -> None:
    rect = (300.0, 600.0, 400.0, 650.0)
    src = draw_pdf(
        tmp_path / "box.pdf", lambda p: p.rect(rect, fill=(0, 0, 0), stroke=None)
    )
    for rotation in (0, 90, 180, 270):
        rotated = edit_pdf(src, tmp_path / f"r{rotation}.pdf", rotate={0: rotation})
        info = {
            "width": 842 if rotation % 180 else 595,
            "height": 595 if rotation % 180 else 842,
            "rotation": rotation,
            "page_width": 595,
            "page_height": 842,
        }
        with pdfium.PdfDocument(rotated) as doc:
            img = doc[0].render(scale=2, rev_byteorder=True).to_numpy()[:, :, :3]
        shown = to_display(rect, info)
        assert np.allclose(_dark_box(img, 2), shown, atol=1)
        assert from_display(shown, info) == rect


def test_bookmarks_and_structure_do_not_survive(tmp_path: Path) -> None:
    src = draw_pdf(tmp_path / "b.pdf", lambda p: p.text(72, 72, "Referto", size=12))
    edit_pdf(src, outline=["Lettera dimissione - ROSSI MARIO"])
    analysis = Pipeline(FakeReader(), FakeNER(())).analyze(src)
    out = tmp_path / "out.pdf"

    release = redact_and_verify(src, analysis, out)

    assert release.report.ok, release.report.issues
    assert not PdfReader(out).outline
    assert b"ROSSI" not in out.read_bytes()
