"""Word export: the anonymized text, line by line, with redacted parts as black blocks.
Layout is preserved only in the PDF; this is for editing the text."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor

from anonmyzation_pipeline.services.anonymization.geometry import contains_point, grow
from anonmyzation_pipeline.services.anonymization.pipeline import Analysis


def export_docx(analysis: Analysis, out: Path) -> None:
    """Write the anonymized text of ``analysis`` to ``out``; redacted parts become black blocks."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(10)
    doc.core_properties.author = ""
    red = {}
    for b in analysis.boxes:
        if b.redact:
            red.setdefault(b.page, []).append(grow(b.rect, 0.5))
    for pno, info in enumerate(analysis.pages):
        layers = [L for L in analysis.layers if L.page == pno]
        layer = (
            next((L for L in layers if L.source == "ocr"), None)
            if info.get("scanned")
            else None
        )
        layer = layer or next((L for L in layers if L.source == "native"), None)
        if pno:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        if layer is None:
            continue
        rects = red.get(pno, [])
        lines: dict[int, list] = {}
        for t in layer.tokens:
            if t.conf < 0.2:
                continue
            lines.setdefault(t.line, []).append(t)
        for toks in lines.values():
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(2)
            for i, t in enumerate(toks):
                if i:
                    para.add_run(" ")
                buf, hidden = "", None
                for ch, cr in zip(t.text, t.chars, strict=False):
                    cx, cy = (cr[0] + cr[2]) / 2, (cr[1] + cr[3]) / 2
                    h = any(contains_point(r, cx, cy) for r in rects)
                    if hidden is not None and h != hidden:
                        _run(para, buf, hidden)
                        buf = ""
                    buf += "█" if h else ch
                    hidden = h
                _run(para, buf, hidden)
    doc.save(str(out))


def _run(para, text: str, hidden: bool | None) -> None:
    if not text:
        return
    r = para.add_run(text)
    if hidden:
        r.font.color.rgb = RGBColor(0, 0, 0)
