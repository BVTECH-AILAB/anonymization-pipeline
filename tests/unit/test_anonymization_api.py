import io
import json
import time
from pathlib import Path

import pypdfium2 as pdfium
import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from anonmyzation_pipeline.errors import (
    FileTooLargeError,
    ServiceUnavailableError,
    TooManyJobsError,
    TooManyPagesError,
)
from anonmyzation_pipeline.services.anonymization.pipeline import Pipeline
from anonmyzation_pipeline.services.anonymization.reader import LiteparseReader
from anonmyzation_pipeline.services.anonymization.verify import Report, verify
from anonmyzation_pipeline.services.jobs import JobManager
from tests.conftest import FAKE_CF
from tests.documents import (
    PATIENT,
    PHONE,
    FakeReader,
    draw_pdf,
    edit_pdf,
    letter_pdf,
    native_words,
    scanned_pdf,
)


def upload(
    client: TestClient,
    path: Path,
    name: str | None = None,
    mime: str = "application/pdf",
    mode: str = "accurate",
) -> dict:
    with path.open("rb") as f:
        response = client.post(
            f"/anonymizations?mode={mode}",
            files={"file": (name or path.name, f, mime)},
        )
    assert response.status_code == 202, response.text
    return response.json()


def wait_done(client: TestClient, job_id: str, timeout: float = 60) -> dict:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        state = client.get(f"/anonymizations/{job_id}").json()
        if state["status"] in ("done", "error", "cancelled"):
            return state
        time.sleep(0.05)
    raise AssertionError("job did not finish")


def pdf_text(data: bytes) -> str:
    with pdfium.PdfDocument(data) as doc:
        return " ".join(
            doc[i].get_textpage().get_text_bounded() for i in range(len(doc))
        )


def test_upload_progress_and_result(client: TestClient, letter: Path) -> None:
    created = upload(client, letter)

    with client.stream("GET", f"/anonymizations/{created['id']}/events") as stream:
        events = [
            json.loads(line[5:])
            for line in stream.iter_lines()
            if line.startswith("data:")
        ]

    assert events[-1]["status"] == "done"
    assert [e["percent"] for e in events] == sorted(e["percent"] for e in events)
    result = events[-1]["result"]
    assert result["verified"] is True
    assert result["counts"]["FISCAL_CODE"] == 1
    assert result["counts"]["PERSON"] >= 2
    assert result["counts"]["QRCODE"] == 1
    assert result["counts"]["SIGNATURE"] == 1
    assert set(result["timings"]) == {
        "text_layers",
        "detection",
        "visual",
        "redaction",
        "verification",
    }


def test_events_of_a_finished_job_end_immediately(
    client: TestClient, letter: Path
) -> None:
    created = upload(client, letter)
    wait_done(client, created["id"])

    with client.stream("GET", f"/anonymizations/{created['id']}/events") as stream:
        lines = [line for line in stream.iter_lines() if line.startswith("data:")]

    assert len(lines) == 1
    assert json.loads(lines[0][5:])["status"] == "done"


def test_download_pdf_without_pii_and_with_generic_name(
    client: TestClient, letter: Path
) -> None:
    created = upload(client, letter, name=f"referto_{PATIENT}_{FAKE_CF}.pdf")
    wait_done(client, created["id"])

    response = client.get(f"/anonymizations/{created['id']}/download/pdf")

    assert response.status_code == 200
    assert "documento_anonimizzato.pdf" in response.headers["content-disposition"]
    assert FAKE_CF not in response.headers["content-disposition"]
    text = pdf_text(response.content)
    for value in (FAKE_CF, PATIENT, PHONE, "Rossi"):
        assert value not in text
    assert "morbo di Crohn" in text
    assert "Babinski" in text
    assert "LETTERA DI DIMISSIONE" in text


def test_download_docx(client: TestClient, letter: Path) -> None:
    created = upload(client, letter)
    wait_done(client, created["id"])

    response = client.get(f"/anonymizations/{created['id']}/download/docx")

    assert response.status_code == 200
    text = " ".join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)
    assert FAKE_CF not in text
    assert "█" in text
    assert "Crohn" in text


def test_pages_with_boxes_are_released_as_images(
    client: TestClient, letter: Path
) -> None:
    created = upload(client, letter)
    state = wait_done(client, created["id"])

    assert state["result"]["verified"] is True
    assert state["result"]["flattened_pages"] == [1]  # 1-based, as shown to the user
    pdf = client.get(f"/anonymizations/{created['id']}/download/pdf").content
    assert PATIENT not in pdf_text(pdf)
    with pdfium.PdfDocument(pdf) as doc:  # rendered page: still a document, not blank
        assert doc[0].render(scale=1).to_numpy().mean() < 255


def test_download_refused_when_a_check_fails(
    client: TestClient, letter: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def failing(*args, **kwargs) -> Report:
        rep = Report()
        rep.fail("value still in text: 'x'")
        return rep

    monkeypatch.setattr(
        "anonmyzation_pipeline.services.anonymization.release.verify", failing
    )
    created = upload(client, letter)
    state = wait_done(client, created["id"])

    assert state["result"]["verified"] is False
    assert state["result"]["issues"]
    response = client.get(f"/anonymizations/{created['id']}/download/pdf")
    assert response.status_code == 409
    assert response.json()["code"] == "VERIFICATION_FAILED"


def test_verification_time_is_measured(
    client: TestClient, letter: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def slow_verify(*args, **kwargs):
        time.sleep(0.05)
        return verify(*args, **kwargs)

    monkeypatch.setattr(
        "anonmyzation_pipeline.services.anonymization.release.verify", slow_verify
    )
    created = upload(client, letter)
    timings = wait_done(client, created["id"])["result"]["timings"]

    assert timings["verification"] >= 0.05


def test_a_failed_edit_leaves_the_job_blocked_and_clean(
    client: TestClient, jobs: JobManager, letter: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    created = upload(client, letter)
    wait_done(client, created["id"])

    def broken(*args, **kwargs):
        raise RuntimeError("redaction failed")

    monkeypatch.setattr(
        "anonmyzation_pipeline.services.anonymization.document.redact_and_verify",
        broken,
    )
    response = client.post(
        f"/anonymizations/{created['id']}/edits", json={"remove": [0]}
    )

    assert response.status_code == 500
    state = client.get(f"/anonymizations/{created['id']}").json()
    assert state["result"]["verified"] is False
    assert state["result"]["edits"] == 0
    assert (
        client.get(f"/anonymizations/{created['id']}/download/pdf").status_code == 409
    )
    assert not list(jobs.get(created["id"]).dir.glob("anonymized.next.*"))


def test_page_previews(client: TestClient, letter: Path) -> None:
    created = upload(client, letter)
    wait_done(client, created["id"])

    for which in ("original", "anonymized"):
        response = client.get(f"/anonymizations/{created['id']}/pages/{which}/0?dpi=60")
        assert response.status_code == 200
        assert response.headers["content-type"] == "image/png"
    assert (
        client.get(f"/anonymizations/{created['id']}/pages/original/3").status_code
        == 404
    )


def test_edits_add_remove_and_search(client: TestClient, letter: Path) -> None:
    created = upload(client, letter)
    result = wait_done(client, created["id"])["result"]
    person = next(b for b in result["boxes"] if b["label"] == "PERSON" and b["redact"])

    edited = client.post(
        f"/anonymizations/{created['id']}/edits",
        json={
            "add": [{"page": 0, "rect": [40, 300, 200, 320]}],
            "remove": [person["id"]],
            "search": "Medicina Interna",
        },
    ).json()

    assert edited["edits"] == 1
    assert edited["search_hits"] == 1
    assert edited["edit_seconds"] >= 0
    boxes = edited["boxes"]
    assert boxes[person["id"]]["redact"] is False
    assert "removed-by-user" in boxes[person["id"]]["sources"]
    assert any(b["sources"] == ["manual-area"] for b in boxes)
    text = pdf_text(client.get(f"/anonymizations/{created['id']}/download/pdf").content)
    assert "Medicina Interna" not in text

    restored = client.post(
        f"/anonymizations/{created['id']}/edits", json={"restore": [person["id"]]}
    ).json()
    assert restored["boxes"][person["id"]]["redact"] is True


def test_edits_before_completion_are_refused(
    client: TestClient, jobs: JobManager, letter: Path
) -> None:
    job = jobs.create(letter.read_bytes(), "a.pdf")
    job.status = "running"

    response = client.post(f"/anonymizations/{job.id}/edits", json={"search": "Rossi"})

    assert response.status_code == 409
    assert response.json()["code"] == "JOB_NOT_READY"


def test_delete_removes_files(
    client: TestClient, jobs: JobManager, letter: Path
) -> None:
    created = upload(client, letter)
    wait_done(client, created["id"])
    job_dir = jobs.get(created["id"]).dir

    assert client.delete(f"/anonymizations/{created['id']}").status_code == 204
    assert not job_dir.exists()
    assert client.get(f"/anonymizations/{created['id']}").status_code == 404


def test_cancel(
    client: TestClient, jobs: JobManager, letter: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    original = jobs.pipeline.analyze

    def slow_analyze(path: Path, progress, visual: bool = True) -> object:
        progress("read", 0)
        time.sleep(0.3)
        progress("read", 0.5)
        return original(path, progress, visual)

    monkeypatch.setattr(jobs.pipeline, "analyze", slow_analyze)
    created = upload(client, letter)
    assert client.post(f"/anonymizations/{created['id']}/cancel").status_code == 202
    deadline = time.monotonic() + 10
    while created["id"] in jobs.jobs and time.monotonic() < deadline:
        time.sleep(0.05)

    assert created["id"] not in jobs.jobs
    assert client.get(f"/anonymizations/{created['id']}").status_code == 404


def test_processing_error_is_reported(
    client: TestClient, jobs: JobManager, letter: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def broken(path: Path, progress, visual: bool = True) -> object:
        raise RuntimeError(f"cannot read {FAKE_CF}")

    monkeypatch.setattr(jobs.pipeline, "analyze", broken)
    created = upload(client, letter)
    state = wait_done(client, created["id"])

    assert state["status"] == "error"
    assert state["error_code"] == "PROCESSING_FAILED"
    assert FAKE_CF not in json.dumps(state)


@pytest.mark.parametrize(
    ("name", "data", "status", "code"),
    [
        ("a.txt", b"hello", 415, "UNSUPPORTED_FORMAT"),
        ("a.pdf", b"not a pdf", 422, "UNREADABLE_DOCUMENT"),
        ("a.png", b"not an image", 422, "UNREADABLE_DOCUMENT"),
    ],
)
def test_upload_errors(
    client: TestClient, name: str, data: bytes, status: int, code: str
) -> None:
    response = client.post(
        "/anonymizations", files={"file": (name, data, "application/octet-stream")}
    )

    assert response.status_code == status
    assert response.json()["code"] == code


def test_password_protected_pdf(client: TestClient, tmp_path: Path) -> None:
    path = draw_pdf(tmp_path / "locked.pdf", lambda p: p.text(72, 72, "segreto"))
    edit_pdf(path, password="pw")

    response = client.post(
        "/anonymizations",
        files={"file": ("locked.pdf", path.read_bytes(), "application/pdf")},
    )

    assert response.status_code == 422
    assert response.json()["code"] == "PASSWORD_PROTECTED"


def test_limits(jobs: JobManager, letter: Path, tmp_path: Path) -> None:
    jobs.settings = jobs.settings.model_copy(update={"max_upload_mb": 1})
    with pytest.raises(FileTooLargeError):
        jobs.create(b"x" * (1024 * 1024 + 1), "big.pdf")

    jobs.settings = jobs.settings.model_copy(
        update={"max_upload_mb": 30, "max_pages": 1}
    )
    many = draw_pdf(tmp_path / "many.pdf", lambda p: None, lambda p: None)
    with pytest.raises(TooManyPagesError):
        jobs.create(many.read_bytes(), "many.pdf")

    jobs.settings = jobs.settings.model_copy(update={"max_pages": 50, "max_jobs": 1})
    jobs.create(letter.read_bytes(), "a.pdf")
    with pytest.raises(TooManyJobsError):
        jobs.create(letter.read_bytes(), "b.pdf")


def test_expired_jobs_are_purged(jobs: JobManager, letter: Path) -> None:
    job = jobs.create(letter.read_bytes(), "a.pdf")
    job.created -= jobs.settings.job_ttl_minutes * 60 + 1

    assert jobs.purge_expired() == 1
    assert job.id not in jobs.jobs
    assert not job.dir.exists()


def test_image_upload_becomes_a_pdf(
    client: TestClient, letter: Path, tmp_path: Path
) -> None:
    with pdfium.PdfDocument(letter) as doc:
        img = doc[0].render(scale=80 / 72).to_pil().convert("RGB")
    buf = io.BytesIO()
    img.save(buf, "PNG")

    response = client.post(
        "/anonymizations", files={"file": ("foto.png", buf.getvalue(), "image/png")}
    )

    assert response.status_code == 202
    state = wait_done(client, response.json()["id"])
    assert state["status"] == "done"
    assert state["result"]["pages"][0]["scanned"] is True


def test_image_with_too_many_frames(jobs: JobManager) -> None:
    jobs.settings = jobs.settings.model_copy(update={"max_pages": 1})
    frames = [Image.new("RGB", (40, 40), c) for c in ("white", "black")]
    buf = io.BytesIO()
    frames[0].save(buf, "TIFF", save_all=True, append_images=frames[1:])

    with pytest.raises(TooManyPagesError):
        jobs.create(buf.getvalue(), "scan.tiff")


def test_letter_fixture_contains_the_pii(letter: Path) -> None:
    text = pdf_text(letter.read_bytes())

    assert FAKE_CF in text
    assert PATIENT in text


def test_scanned_upload_uses_ocr(
    client: TestClient, fake_reader: FakeReader, tmp_path: Path
) -> None:
    native = letter_pdf(tmp_path / "native.pdf")
    scan = scanned_pdf(native, tmp_path / "scan.pdf")
    fake_reader.words = native_words(native)

    created = upload(client, scan)
    result = wait_done(client, created["id"])["result"]

    assert result["verified"] is True
    assert result["pages"][0]["scanned"] is True
    assert result["counts"]["FISCAL_CODE"] == 1
    assert all(b["layer"] in ("ocr", "visual") for b in result["boxes"])


def test_upload_refused_while_models_load(
    settings, tmp_path: Path, letter: Path
) -> None:
    loading = JobManager(
        Pipeline(LiteparseReader(""), None),  # no OCR server: never ready
        settings.model_copy(update={"preload_models": True}),
        tmp_path / "w",
    )
    try:
        with pytest.raises(ServiceUnavailableError):
            loading.create(letter.read_bytes(), "a.pdf")
    finally:
        loading.shutdown()


def test_mode_defaults_to_accurate(client: TestClient, letter: Path) -> None:
    with letter.open("rb") as f:
        created = client.post(
            "/anonymizations", files={"file": ("a.pdf", f, "application/pdf")}
        ).json()

    assert created["mode"] == "accurate"
    assert client.get(f"/anonymizations/{created['id']}").json()["mode"] == "accurate"


def test_unknown_mode_is_refused(client: TestClient, letter: Path) -> None:
    with letter.open("rb") as f:
        response = client.post(
            "/anonymizations?mode=turbo",
            files={"file": ("a.pdf", f, "application/pdf")},
        )

    assert response.status_code == 422


def test_fast_mode_releases_anonymized_markdown(
    client: TestClient, letter: Path
) -> None:
    created = upload(client, letter, name=f"referto_{FAKE_CF}.pdf", mode="fast")
    state = wait_done(client, created["id"])

    assert created["mode"] == "fast"
    assert state["mode"] == "fast"
    result = state["result"]
    assert result["verified"] is True
    assert result["counts"]["FISCAL_CODE"] == 1
    assert "boxes" not in result
    assert "QRCODE" not in result["counts"]  # no visual detection in fast mode
    response = client.get(f"/anonymizations/{created['id']}/download/md")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/markdown")
    assert "documento_anonimizzato.md" in response.headers["content-disposition"]
    md = response.text
    assert md.startswith("<!-- pagina 1 -->")
    assert "[CODICE FISCALE]" in md
    assert "[IMMAGINE]" in md
    for value in (FAKE_CF, PATIENT, PHONE):
        assert value not in md
    assert "Crohn" in md


def test_fast_markdown_refused_when_a_value_survives(
    client: TestClient, letter: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        "anonmyzation_pipeline.services.anonymization.document.verify_markdown",
        lambda *_: ["value still in the Markdown export: 'x'"],
    )
    created = upload(client, letter, mode="fast")
    state = wait_done(client, created["id"])

    assert state["result"]["verified"] is False
    response = client.get(f"/anonymizations/{created['id']}/download/md")
    assert response.status_code == 409
    assert response.json()["code"] == "VERIFICATION_FAILED"


@pytest.mark.parametrize(
    ("mode", "method", "path"),
    [
        ("fast", "get", "download/pdf"),
        ("fast", "get", "download/docx"),
        ("fast", "get", "preview"),
        ("fast", "get", "pages/anonymized/0"),
        ("fast", "post", "edits"),
        ("accurate", "get", "download/md"),
    ],
)
def test_actions_of_the_other_mode_are_refused(
    client: TestClient, letter: Path, mode: str, method: str, path: str
) -> None:
    created = upload(client, letter, mode=mode)
    wait_done(client, created["id"])

    url = f"/anonymizations/{created['id']}/{path}"
    response = client.post(url, json={}) if method == "post" else client.get(url)

    assert response.status_code == 409
    assert response.json()["code"] == "MODE_NOT_SUPPORTED"


def test_fast_job_keeps_the_original_preview(client: TestClient, letter: Path) -> None:
    created = upload(client, letter, mode="fast")
    wait_done(client, created["id"])

    response = client.get(f"/anonymizations/{created['id']}/pages/original/0?dpi=60")

    assert response.status_code == 200


def test_preview_manifest(client: TestClient, letter: Path) -> None:
    created = upload(client, letter)
    wait_done(client, created["id"])
    url = f"/anonymizations/{created['id']}/preview"

    preview = client.get(url).json()

    assert preview["id"] == created["id"]
    assert preview["mode"] == "accurate"
    assert preview["version"] == 0
    assert preview["verified"] is True
    page = preview["pages"][0]
    assert page["index"] == 0
    assert (page["width"], page["height"]) == (595, 842)
    assert page["original"] == f"/anonymizations/{created['id']}/pages/original/0"
    assert page["anonymized"] == (
        f"/anonymizations/{created['id']}/pages/anonymized/0?v=0"
    )
    assert client.get(page["original"]).status_code == 200
    assert client.get(page["anonymized"]).status_code == 200
    box = preview["boxes"][0]
    assert set(box) == {"id", "page", "rect", "label", "layer", "redact"}
    assert preview["labels"]["PERSON"] == "Nomi"
    assert FAKE_CF not in json.dumps(preview)
    assert PATIENT not in json.dumps(preview)

    client.post(f"/anonymizations/{created['id']}/edits", json={"remove": [0]})
    edited = client.get(url).json()

    assert edited["version"] == 1
    assert edited["pages"][0]["anonymized"].endswith("?v=1")
    assert edited["boxes"][0]["redact"] is False


def test_preview_before_completion_is_refused(
    client: TestClient, jobs: JobManager, letter: Path
) -> None:
    job = jobs.create(letter.read_bytes(), "a.pdf")
    job.status = "running"

    response = client.get(f"/anonymizations/{job.id}/preview")

    assert response.status_code == 409
    assert response.json()["code"] == "JOB_NOT_READY"
